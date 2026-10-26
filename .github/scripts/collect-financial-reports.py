#!/usr/bin/env python3
"""Collect and archive VietstockFinance statements first; OCR is handled by the separate post-archive workflow."""
from __future__ import annotations
import hashlib, json, os, re, sys, time, io, zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse, unquote, parse_qs
import requests
from bs4 import BeautifulSoup

ROOT = Path(__file__).resolve().parents[2]
QUEUE_PATH = ROOT / "config/report-queue.json"
SEC_PATH = ROOT / "config/securities.json"
SOURCES_PATH = ROOT / "config/source-registry.json"
OUT = ROOT / "reports/raw"
MAX_BYTES = int(os.getenv("MAX_REPORT_BYTES", str(75 * 1024 * 1024)))
MAX_PER_JOB = min(max(1, int(os.getenv("MAX_DOWNLOADS_PER_JOB", "2"))), 2)
MAX_JOBS_PER_RUN = 1  # Strictly one queued ticker job per run, regardless of inputs.
USER_AGENT = "solieutaichinh-report-collector/1.0 (public financial report archive)"
SESSION = requests.Session()
SESSION.headers.update({"User-Agent": USER_AGENT, "Accept": "text/html,application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.8"})
NOW = datetime.now(timezone.utc).isoformat()

def load(path, fallback):
    if not path.exists(): return fallback
    return json.loads(path.read_text(encoding="utf-8"))

def save(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

def domain_allowed(url, allowed):
    host = (urlparse(url).hostname or "").lower().rstrip(".")
    return any(host == d or host.endswith("." + d) for d in allowed)

def search_web(query):
    url = "https://html.duckduckgo.com/html/"
    try:
        response = SESSION.post(url, data={"q": query}, timeout=12)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        found = []
        for a in soup.select("a.result__a"):
            href = a.get("href", "")
            title = a.get_text(" ", strip=True)
            parsed = urlparse(href)
            if "duckduckgo.com" in (parsed.hostname or ""):
                target = parse_qs(parsed.query).get("uddg", [""])[0]
                if target.startswith("http"):
                    href = target
            if href.startswith("http"):
                found.append({"url": href, "title": title})
        return found
    except Exception as exc:
        print(json.dumps({"event":"search_failed","query":query,"error":str(exc)[:500]}, ensure_ascii=False))
        return [{"error": str(exc), "query": query}]

def extract_file(path, content_type):
    suffix = path.suffix.lower()
    if suffix == ".pdf" or "pdf" in content_type.lower():
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            pages = []
            for i, page in enumerate(reader.pages[:20]):
                text = page.extract_text() or ""
                if text.strip(): pages.append({"page": i + 1, "text": text[:25000]})
            return {"kind":"pdf", "pages":pages, "page_count":len(reader.pages), "extraction":"text", "note":None if pages else "PDF không có lớp văn bản; có thể là bản scan cần OCR thủ công."}
        except Exception as exc:
            return {"kind":"pdf", "pages":[], "extraction":"failed", "note":str(exc)}
    if suffix == ".zip" or "zip" in content_type.lower():
        try:
            from pypdf import PdfReader
            entries = []
            with zipfile.ZipFile(path, "r") as archive:
                for member in archive.infolist():
                    if member.is_dir():
                        continue
                    name = member.filename.replace("\\", "/")
                    # Ignore path traversal, nested archives and non-report payloads.
                    if name.startswith("/") or any(part == ".." for part in name.split("/")):
                        continue
                    if not name.lower().endswith((".pdf", ".xlsx", ".xls")) or member.file_size > MAX_BYTES:
                        continue
                    raw = archive.read(member)
                    if name.lower().endswith(".pdf"):
                        try:
                            reader = PdfReader(io.BytesIO(raw))
                            pages = []
                            for i, page in enumerate(reader.pages[:20]):
                                page_text = page.extract_text() or ""
                                if page_text.strip():
                                    pages.append({"page":i+1,"text":page_text[:25000]})
                            entries.append({"name":name,"kind":"pdf","bytes":len(raw),
                                            "pages":pages,"page_count":len(reader.pages),
                                            "extraction":"text",
                                            "note":None if pages else "PDF trong ZIP không có lớp văn bản; cần xem thủ công."})
                        except Exception as exc:
                            entries.append({"name":name,"kind":"pdf","bytes":len(raw),
                                            "pages":[],"extraction":"failed","note":str(exc)[:500]})
                    else:
                        entries.append({"name":name,"kind":"spreadsheet","bytes":len(raw),
                                        "extraction":"archive_member_not_parsed"})
            return {"kind":"zip","members":entries,"member_count":len(entries),
                    "extraction":"archive_members","note":None if entries else "ZIP không có PDF/Excel báo cáo hỗ trợ."}
        except Exception as exc:
            return {"kind":"zip","members":[],"extraction":"failed","note":str(exc)[:1000]}
    if suffix == ".xlsx" or "spreadsheetml" in content_type.lower():
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for ws in wb.worksheets[:30]:
                rows = []
                for row in ws.iter_rows(min_row=1, max_row=250, values_only=True):
                    vals = [str(v)[:500] if v is not None else None for v in row[:40]]
                    if any(v is not None for v in vals): rows.append(vals)
                sheets.append({"name":ws.title,"rows":rows})
            return {"kind":"xlsx","sheets":sheets,"extraction":"table"}
        except Exception as exc:
            return {"kind":"xlsx","sheets":[],"extraction":"failed","note":str(exc)}
    return {"kind":"unknown","extraction":"not_supported","note":"Tệp được lưu nhưng định dạng chưa hỗ trợ trích xuất tự động."}

def safe_name(url, content_type):
    path = unquote(urlparse(url).path)
    suffix = Path(path).suffix.lower()
    if suffix not in {".pdf",".xlsx",".xls",".zip"}:
        suffix = ".pdf" if "pdf" in content_type.lower() else ".xlsx" if "spreadsheetml" in content_type.lower() else ".zip" if "zip" in content_type.lower() else ".bin"
    return suffix

def is_official_statement_title(title, symbol, require_symbol=True):
    """Reject analyst/research material; accept only issuer filing notices for actual statements."""
    t = (title or "").lower()
    sym = re.escape(symbol.lower())
    if require_symbol and not re.search(rf"\b{sym}\b", t):
        return False
    reject = [
        "phân tích", "bao cao phan tich", "báo cáo phân tích", "cập nhật kqkd",
        "khuyến nghị", "định giá", "price target", "analyst", "earning report",
        "investment thesis", "forecast", "dự phóng", "triển vọng", "update kqkd",
        "mua", "bán", "nắm giữ", "brokerage"
    ]
    if any(word in t for word in reject):
        return False
    return any(word in t for word in [
        "báo cáo tài chính", "bao cao tai chinh", "financial statements",
        "financial statement", "consolidated financial statements"
    ])

def is_vietstock_statement_pdf(url, symbol):
    """Recognize direct Vietstock BCTC PDFs by their URL path, even if search metadata is poor."""
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower()
    path = unquote(parsed.path).lower()
    if host not in {"static1.vietstock.vn", "static2.vietstock.vn"}:
        return False
    if not path.endswith(".pdf") or "/bctc/" not in path or "/vn/" not in path:
        return False
    # Typical filename: MBB_Baocaotaichinh_6T_2026_Soatxet_Hopnhat.pdf
    filename = Path(path).name
    if not re.match(rf"^{re.escape(symbol.lower())}[_-]", filename):
        return False
    return any(token in filename for token in ("baocaotaichinh", "bao-cao-tai-chinh", "bctc", "financialstatement", "financial-statement"))

def expand_vietstock_filing_page(candidate, symbol):
    """Vietstock is used only as an index of issuer disclosure notices, not as an analyst source."""
    url = candidate.get("url", "")
    host = (urlparse(url).hostname or "").lower()
    if host not in {"vietstock.vn", "www.vietstock.vn", "en.vietstock.vn"}:
        return []
    page_title = candidate.get("title", "")
    if not is_official_statement_title(page_title, symbol):
        return []
    try:
        response = SESSION.get(url, timeout=25, allow_redirects=True)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        heading = soup.find(["h1", "title"])
        verified_title = heading.get_text(" ", strip=True) if heading else page_title
        if not is_official_statement_title(verified_title, symbol):
            return []
        attachments = []
        for a in soup.find_all("a", href=True):
            href = a["href"]
            label = a.get_text(" ", strip=True)
            absolute = requests.compat.urljoin(response.url, href)
            path = urlparse(absolute).path.lower()
            if not path.endswith((".pdf", ".xlsx", ".xls")):
                continue
            if not domain_allowed(absolute, ["static1.vietstock.vn", "static2.vietstock.vn", "vietstock.vn"]):
                continue
            evidence = (label + " " + path).lower()
            if any(x in evidence for x in ["explanation", "giai-trinh", "giai_trinh", "forecast", "phan-tich", "analyst"]):
                continue
            if not any(x in evidence for x in ["baocao", "bao-cao", "financial", "bctc", "statement", "tai-chinh"]):
                continue
            attachments.append({
                "url": absolute, "title": label or verified_title,
                "host": (urlparse(absolute).hostname or "").lower(),
                "allowlisted": True, "official_filing_page": True,
                "filing_page": response.url, "filing_page_title": verified_title
            })
        return attachments
    except Exception:
        return []

def vietstock_document_index_url(symbol):
    """The per-symbol Vietstock document index is the primary discovery source."""
    return f"https://finance.vietstock.vn/{symbol}/tai-tai-lieu.htm?doctype=1"

def is_financial_statement_document(text):
    """Keep BCTC documents and reject analyst research, estimates and recommendations."""
    t = (text or "").casefold()
    reject = [
        "báo cáo phân tích", "bao cao phan tich", "phân tích cổ phiếu",
        "khuyến nghị", "định giá", "giá mục tiêu", "dự phóng", "du phong",
        "forecast", "triển vọng", "cập nhật kqkd", "equity research",
        "analyst", "brokerage", "mua", "bán", "nắm giữ", "annual report",
        "khóa học", "khoa hoc", "đào tạo", "dao tao", "training", "đọc hiểu", "doc hieu", "webinar", "seminar"
    ]
    if any(term in t for term in reject):
        return False
    return any(term in t for term in [
        "báo cáo tài chính", "bao cao tai chinh", "bctc",
        "financial statement", "financial statements"
    ])

def _document_context(anchor):
    """Get the closest useful row/card around a document link."""
    for tag in ("tr", "li", "article"):
        node = anchor.find_parent(tag)
        if node and len(node.get_text(" ", strip=True)) > 20:
            return node
    node = anchor.parent
    for _ in range(3):
        if not node:
            break
        text = node.get_text(" ", strip=True)
        if len(text) > 20:
            return node
        node = node.parent
    return anchor

def _extract_report_year_and_date(text):
    text = text or ""
    year_matches = [int(x) for x in re.findall(r"\b(20\d{2})\b", text)]
    year = max(year_matches) if year_matches else None
    date = None
    m = re.search(r"\b(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})\b", text)
    if m:
        try:
            date = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3))).date().isoformat()
        except ValueError:
            pass
    if not date:
        m = re.search(r"\b(\d{1,2})[-/.](\d{1,2})[-/.](20\d{2})\b", text)
        if m:
            try:
                date = datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date().isoformat()
            except ValueError:
                pass
    return year, date

def _vietstock_epoch_date(value):
    """Convert Vietstock /Date(milliseconds)/ metadata into a stable ISO date."""
    match = re.search(r"/Date\((-?\d+)\)/", str(value or ""))
    if not match:
        return None
    try:
        stamp = int(match.group(1)) / 1000.0
        return datetime.fromtimestamp(stamp, tz=timezone.utc).date().isoformat()
    except (OverflowError, OSError, ValueError):
        return None

def _period_from_report_text(text):
    t = (text or "").casefold()
    m = re.search(r"\b(?:quý|q)\s*([1-4i]{1,2})\D{0,16}(20\d{2})", t)
    if m:
        q = m.group(1).upper()
        q = {"I":"1","II":"2","III":"3","IV":"4"}.get(q, q)
        if q.isdigit() and int(q) in range(1,5):
            return f"{m.group(2)}-Q{q}"
    m = re.search(r"\b(20\d{2})\D{0,12}(?:quý|q)\s*([1-4])\b", t)
    if m:
        return f"{m.group(1)}-Q{m.group(2)}"
    m = re.search(r"\b(?:6 tháng|bán niên|6m|6t)\D{0,12}(20\d{2})", t)
    if m:
        return f"{m.group(1)}-H1"
    m = re.search(r"\b(?:9 tháng|9m|9t)\D{0,12}(20\d{2})", t)
    if m:
        return f"{m.group(1)}-9M"
    m = re.search(r"\b(?:năm|year)\D{0,8}(20\d{2})", t)
    if m:
        return f"{m.group(1)}-FY"
    return None

def _urls_from_document_page(url, symbol, row_text, index_url):
    """Expand a Vietstock document detail page to its direct PDF/XLS attachments."""
    host = (urlparse(url).hostname or "").lower()
    path = urlparse(url).path.lower()
    if not (host == "finance.vietstock.vn" and path.startswith("/downloadedoc/")):
        return [{"url":url, "title":row_text[:300], "symbol":symbol, "allowlisted":True,
                 "query":"Vietstock document index", "host":host, "page_url":index_url,
                 "document_index":index_url}]
    try:
        response = SESSION.get(url, timeout=30, allow_redirects=True)
        response.raise_for_status()
        ctype = (response.headers.get("Content-Type") or "").lower()
        if "pdf" in ctype or response.content[:5] == b"%PDF-":
            return [{"url":response.url, "title":row_text[:300], "symbol":symbol,
                     "allowlisted":True, "query":"Vietstock document index", "host":urlparse(response.url).hostname,
                     "page_url":index_url, "document_index":index_url}]
        soup = BeautifulSoup(response.text, "html.parser")
        found = []
        for a in soup.find_all("a", href=True):
            href = requests.compat.urljoin(response.url, a["href"])
            ap = urlparse(href)
            label = a.get_text(" ", strip=True)
            evidence = (row_text + " " + label + " " + href).casefold()
            if ap.hostname and (ap.hostname.lower().endswith("vietstock.vn")) and (
                ap.path.lower().endswith((".pdf", ".xlsx", ".xls")) or
                (ap.hostname.lower() in {"static1.vietstock.vn","static2.vietstock.vn"} and ".pdf" in ap.path.lower())
            ):
                if is_financial_statement_document(evidence) or is_financial_statement_document(row_text):
                    found.append({"url":href,"title":label or row_text[:300],"symbol":symbol,
                                  "allowlisted":True,"query":"Vietstock document index",
                                  "host":ap.hostname.lower(),"page_url":response.url,
                                  "document_index":index_url,"document_detail_url":response.url})
        # Some detail pages embed attachment URLs inside script/json instead of <a> elements.
        for href in re.findall(r"""https?://[^"'\s<>]+?\.(?:pdf|xlsx?)(?:\?[^"'\s<>]*)?""", response.text, re.I):
            ap = urlparse(href)
            evidence = (row_text + " " + href).casefold()
            if ap.hostname and ap.hostname.lower().endswith("vietstock.vn") and (
                is_financial_statement_document(evidence) or is_financial_statement_document(row_text)
            ):
                found.append({"url":href,"title":row_text[:300],"symbol":symbol,
                              "allowlisted":True,"query":"Vietstock document index",
                              "host":ap.hostname.lower(),"page_url":response.url,
                              "document_index":index_url,"document_detail_url":response.url})
        # If Vietstock serves the BCTC directly at /downloadedoc/{id}, let the
        # normal downloader validate Content-Type and PDF signature rather than guessing.
        if not found and is_financial_statement_document(row_text):
            found.append({"url":response.url,"title":row_text[:300],"symbol":symbol,
                          "allowlisted":True,"query":"Vietstock document index",
                          "host":urlparse(response.url).hostname,"page_url":response.url,
                          "document_index":index_url})
        return found
    except Exception as exc:
        print(json.dumps({"event":"vietstock_document_detail_failed","symbol":symbol,
                          "url":url,"error":str(exc)[:500]},ensure_ascii=False))
        return []

def discover_for_job(job, allowed_domains):
    """Use only VietstockFinance's own AJAX-backed per-symbol document list."""
    symbol = str(job["symbol"]).upper()
    start_year = int(job.get("startYear", 2024))
    index_url = vietstock_document_index_url(symbol)
    base_url = "https://finance.vietstock.vn"
    requested_type = str(job.get("reportType", "both")).casefold()

    # Targeted single-document repair jobs must not query every year/page of
    # the document catalog. Validate and return only their explicitly supplied
    # VietstockFinance URL, so one-ticker repairs stay fast and predictable.
    requested_urls = job.get("requestedUrls")
    if isinstance(requested_urls, list) and requested_urls:
        targeted = []
        for request in requested_urls:
            if not isinstance(request, dict):
                continue
            requested_url = str(request.get("url") or "").strip()
            requested_title = str(request.get("title") or "").strip()
            parsed = urlparse(requested_url)
            hostname = (parsed.hostname or "").lower()
            path = unquote(parsed.path).lower()
            filename = Path(unquote(parsed.path)).name.lower()
            if parsed.scheme != "https" or hostname not in {"static1.vietstock.vn", "static2.vietstock.vn"}:
                continue
            if not filename.startswith(symbol.lower() + "_") or not filename.endswith(".pdf"):
                continue
            if "/bctc/" not in path or "/vn/" not in path:
                continue
            if not is_financial_statement_document(requested_title):
                continue
            if not re.search(r"hợp\s*nhất|hop\s*nhat|consolidated", requested_title, re.I):
                continue
            if re.search(r"công\s*ty\s*mẹ|cong\s*ty\s*me|riêng\s*lẻ|rieng\s*le|parent\s+company|separate", requested_title, re.I):
                continue
            targeted.append({
                "url": requested_url,
                "title": requested_title,
                "host": hostname,
                "allowlisted": domain_allowed(requested_url, allowed_domains),
                "symbol": symbol,
                "query": "VietstockFinance targeted source repair",
                "page_url": index_url,
                "document_index": index_url,
                "document_id": request.get("document_id"),
                "period": str(request.get("period") or ""),
                "report_year": int(request.get("report_year") or start_year),
                "published_at": request.get("published_at"),
                "report_type": "consolidated",
                "vietstock_title": requested_title,
                "full_name": requested_title,
                "file_ext": ".pdf",
            })
        print(json.dumps({
            "event": "vietstock_targeted_source_repair",
            "symbol": symbol,
            "requested_url_count": len(requested_urls),
            "accepted_url_count": len(targeted),
            "periods": [item.get("period") for item in targeted],
        }, ensure_ascii=False))
        return targeted

    try:
        page_response = SESSION.get(index_url, timeout=35, allow_redirects=True)
        page_response.raise_for_status()
    except Exception as exc:
        print(json.dumps({"event":"vietstock_index_fetch_failed","symbol":symbol,
                          "url":index_url,"error":str(exc)[:800]},ensure_ascii=False))
        return []

    soup = BeautifulSoup(page_response.text, "html.parser")
    page_title = soup.title.get_text(" ",strip=True) if soup.title else ""
    # Match the site's own request token: hidden input first, then JS assignment.
    token = ""
    token_input = soup.find("input", attrs={"name":"__RequestVerificationToken"})
    if token_input:
        token = str(token_input.get("value") or "")
    if not token:
        match = re.search(r"""(?:_tokenString|__RequestVerificationToken)\s*[:=]\s*['"]([^'"]+)['"]""",
                          page_response.text, re.I)
        if match:
            token = match.group(1)
    try:
        doctype = int(parse_qs(urlparse(index_url).query).get("doctype", ["1"])[0] or 1)
    except (TypeError, ValueError):
        doctype = 1
    headers = {
        "Referer": page_response.url,
        "Origin": base_url,
        "X-Requested-With": "XMLHttpRequest",
        "Accept": "application/json, text/javascript, */*; q=0.01",
    }

    def call_json(path, method="GET", payload=None):
        url = requests.compat.urljoin(base_url, path)
        try:
            if method == "POST":
                data = dict(payload or {})
                if token:
                    data["__RequestVerificationToken"] = token
                resp = SESSION.post(url, data=data, headers=headers, timeout=15)
            else:
                resp = SESSION.get(url, headers=headers, timeout=15)
            ctype = resp.headers.get("Content-Type","")
            try:
                parsed = resp.json()
            except Exception:
                parsed = None
            print(json.dumps({
                "event":"vietstock_api_probe","symbol":symbol,"endpoint":path,
                "status":resp.status_code,"content_type":ctype,
                "bytes":len(resp.content),"json_type":type(parsed).__name__ if parsed is not None else None,
                "json_preview":str(parsed)[:700] if parsed is not None else resp.text[:500]
            },ensure_ascii=False))
            resp.raise_for_status()
            return parsed
        except Exception as exc:
            print(json.dumps({"event":"vietstock_api_failed","symbol":symbol,
                              "endpoint":path,"error":str(exc)[:600]},ensure_ascii=False))
            return None

    # Same endpoints as the public page, but request only years in the user's scope.
    call_json("/data/getdocumenttype", "GET")
    year_response = call_json("/data/getyearreport", "POST", {"code":symbol})

    def years_from_response(value):
        found = set()
        def walk(node):
            if isinstance(node, dict):
                for key, item in node.items():
                    if str(key).casefold() in {"year","reportyear","value","id"}:
                        if isinstance(item, (int,float)) or (isinstance(item,str) and re.fullmatch(r"20\d{2}",item.strip())):
                            y = int(item)
                            if start_year <= y <= datetime.now(timezone.utc).year + 1:
                                found.add(y)
                    walk(item)
            elif isinstance(node, list):
                for item in node:
                    walk(item)
            elif isinstance(node, (int,float)) or (isinstance(node,str) and re.fullmatch(r"20\d{2}",node.strip())):
                y = int(node)
                if start_year <= y <= datetime.now(timezone.utc).year + 1:
                    found.add(y)
        walk(value)
        return found

    # The endpoint usually returns the years available on this symbol's Vietstock
    # index. Add each calendar year through the current year as a safe fallback.
    years_to_query = set(range(start_year, datetime.now(timezone.utc).year + 1))
    years_to_query.update(years_from_response(year_response))
    years_to_query = sorted(years_to_query)
    def rows_from_json(value):
        if isinstance(value, list):
            return value
        if isinstance(value, dict):
            for key in ("data","Data","items","Items","result","Result","rows","Rows"):
                sub = value.get(key)
                if isinstance(sub, list):
                    return sub
        return []

    seen_rows = set()
    all_rows = []
    total_rows_by_year = {}
    def add_rows(items):
        for row in items:
            if not isinstance(row, dict):
                continue
            ident = str(row.get("Id") or row.get("FileInfoID") or row.get("Url") or row.get("FullName") or row.get("Title") or "")
            if ident and ident not in seen_rows:
                seen_rows.add(ident)
                all_rows.append(row)

    for report_year in years_to_query:
        first_page = call_json("/data/getdocument", "POST",
                               {"code":symbol,"page":1,"type":doctype,"year":report_year})
        rows = rows_from_json(first_page)
        add_rows(rows)
        total_rows = 0
        if rows:
            try:
                total_rows = int(rows[0].get("TotalRow") or rows[0].get("totalRow") or 0)
            except (TypeError, ValueError):
                total_rows = 0
        total_rows_by_year[str(report_year)] = total_rows or len(rows)
        # Bound the catalog scan during a single-ticker run. Eight pages means
        # at most 160 records per year; this is a broad margin for quarterly
        # financial statements while avoiding dozens of slow requests if the
        # site's TotalRow metadata is unexpectedly high. Older/missing periods
        # can still be added later via targeted source URLs.
        max_pages_per_year = 8
        if total_rows:
            pages = min(max((total_rows + 19) // 20, 1), max_pages_per_year)
        else:
            pages = 1 if len(rows) < 20 else max_pages_per_year
        for page in range(2, pages + 1):
            result = call_json("/data/getdocument", "POST",
                               {"code":symbol,"page":page,"type":doctype,"year":report_year})
            page_rows = rows_from_json(result)
            if not page_rows:
                break
            before = len(all_rows)
            add_rows(page_rows)
            if len(all_rows) == before or len(page_rows) < 20:
                break
    candidates = []
    for row in all_rows:
        title = str(row.get("FullName") or row.get("Title") or row.get("FileName") or "").strip()
        # Some rows use separate Title/FullName/Url labels; combine metadata for filtering.
        name = str(row.get("Title") or "")
        full_name = str(row.get("FullName") or "")
        file_url = str(row.get("Url") or row.get("FileUrl") or row.get("DownloadUrl") or "").strip()
        updated = str(row.get("LastUpdate") or row.get("PublishDate") or row.get("PublishedAt") or "")
        evidence = " ".join([title,name,full_name,file_url,updated])
        if not is_financial_statement_document(evidence):
            continue
        if any(word in evidence.casefold() for word in [
            "báo cáo phân tích","bao cao phan tich","khuyến nghị","định giá",
            "dự phóng","forecast","equity research","brokerage","triển vọng",
            "khóa học","khoa hoc","đào tạo","dao tao","training","webinar","seminar"
        ]):
            continue
        period = _period_from_report_text(evidence)
        if job.get("quarterly_only") and not re.fullmatch(r"20\d{2}-Q[1-4]", str(period or "")):
            continue
        year, published_at = _extract_report_year_and_date(evidence)
        period_year = re.match(r"(20\d{2})-", period or "")
        if period_year:
            year = int(period_year.group(1))
        if year is not None and year < start_year:
            continue
        if year is None:
            continue
        if requested_type == "consolidated" and re.search(r"riêng\s*lẻ|rieng\s*le|separate financial|công ty\s*mẹ|cong ty\s*me|parent company", evidence, re.I):
            continue
        if requested_type == "separate" and re.search(r"hợp\s*nhất|hop\s*nhat|consolidated", evidence, re.I):
            continue
        if not file_url:
            continue
        absolute = requests.compat.urljoin(base_url, file_url)
        host = (urlparse(absolute).hostname or "").lower()
        if not host.endswith("vietstock.vn") and host != "vietstock.vn":
            continue
        candidates.append({
            "url":absolute,"title":title or name or full_name or absolute,
            "host":host,"allowlisted":domain_allowed(absolute,allowed_domains),
            "symbol":symbol,"query":"VietstockFinance /data/getdocument",
            "page_url":index_url,"document_index":index_url,
            "document_id":row.get("FileInfoID") or row.get("Id"),
            "report_year":year,"published_at":published_at or _vietstock_epoch_date(updated),
            "period":period,"report_type":job.get("reportType","both"),
            "vietstock_title":name,"full_name":full_name,
            "file_ext":row.get("FileExt"),"last_update":updated,
            "total_rows":total_rows
        })
    requested_urls = job.get("requestedUrls")
    if isinstance(requested_urls, list) and requested_urls:
        targeted = []
        for request in requested_urls:
            if not isinstance(request, dict):
                continue
            requested_url = str(request.get("url") or "").strip()
            requested_title = str(request.get("title") or "").strip()
            parsed_url = urlparse(requested_url)
            file_name = Path(unquote(parsed_url.path)).name.lower()
            if parsed_url.scheme != "https" or (parsed_url.hostname or "").lower() not in {
                "static1.vietstock.vn", "static2.vietstock.vn"
            }:
                continue
            if not file_name.startswith(symbol.lower() + "_") or not file_name.endswith(".pdf"):
                continue
            if not is_financial_statement_document(requested_title):
                continue
            if not re.search(r"hợp\s*nhất|hop\s*nhat|consolidated", requested_title, re.I):
                continue
            if re.search(r"công\s*ty\s*mẹ|cong\s*ty\s*me|riêng\s*lẻ|rieng\s*le|parent\s+company|separate", requested_title, re.I):
                continue
            matching = next((candidate for candidate in candidates
                             if canonical_document_url(candidate.get("url")) == canonical_document_url(requested_url)), None)
            candidate = dict(matching or {})
            candidate.update({
                "url": requested_url,
                "title": requested_title,
                "host": parsed_url.hostname,
                "allowlisted": domain_allowed(requested_url, allowed_domains),
                "symbol": symbol,
                "query": "VietstockFinance targeted source repair",
                "page_url": index_url,
                "document_index": index_url,
                "period": str(request.get("period") or candidate.get("period") or ""),
                "report_year": int(request.get("report_year") or candidate.get("report_year") or start_year),
                "published_at": request.get("published_at") or candidate.get("published_at"),
                "report_type": "consolidated",
                "vietstock_title": requested_title,
                "full_name": requested_title,
                "file_ext": ".pdf",
            })
            targeted.append(candidate)
        # Targeted jobs never fall back to the broad index result set.
        candidates = targeted

    print(json.dumps({
        "event":"vietstock_index_discovery","symbol":symbol,"index_url":index_url,
        "http_status":page_response.status_code,"page_title":page_title,
        "ajax_rows":len(all_rows),"total_rows":total_rows,
        "candidate_count":len(candidates),
        "sample_rows":[{k:r.get(k) for k in ("Title","FullName","Url","FileExt","LastUpdate","TotalRow")} for r in all_rows[:5]]
    },ensure_ascii=False))
    deduped=[]
    seen_urls=set()
    for candidate in candidates:
        if candidate["url"] not in seen_urls:
            seen_urls.add(candidate["url"])
            deduped.append(candidate)
    return deduped

def download_candidate(candidate, job_dir):
    url = candidate["url"]
    host = (urlparse(url).hostname or "").lower()
    approved_attachment = candidate.get("official_filing_page") is True and host in {"static1.vietstock.vn", "static2.vietstock.vn", "vietstock.vn"}
    if not candidate.get("allowlisted") or (not domain_allowed(url, ALLOWED_DOMAINS) and not approved_attachment):
        return {"url":url,"title":candidate.get("title"),"status":"blocked_untrusted_domain","host":host}
    try:
        with SESSION.get(url, timeout=18, stream=True, allow_redirects=True) as response:
            response.raise_for_status()
            final_url = response.url
            final_host = (urlparse(final_url).hostname or "").lower()
            redirect_is_filing_attachment = (
                (candidate.get("official_filing_page") is True and final_host in {"static1.vietstock.vn", "static2.vietstock.vn", "vietstock.vn"})
                or is_vietstock_statement_pdf(final_url, candidate.get("symbol", ""))
            )
            if not domain_allowed(final_url, ALLOWED_DOMAINS) and not redirect_is_filing_attachment:
                return {"url":url,"status":"blocked_redirect_outside_allowlist","final_url":final_url}
            ctype = response.headers.get("Content-Type","")
            if not ("pdf" in ctype.lower() or "spreadsheet" in ctype.lower() or "zip" in ctype.lower() or urlparse(final_url).path.lower().endswith((".pdf",".xlsx",".xls",".zip"))):
                return {"url":url,"status":"skipped_not_document","content_type":ctype}
            declared = int(response.headers.get("Content-Length","0") or 0)
            if declared > MAX_BYTES: return {"url":url,"status":"skipped_too_large","bytes":declared}
            data = bytearray()
            for chunk in response.iter_content(1024*64):
                if not chunk: continue
                data.extend(chunk)
                if len(data) > MAX_BYTES: return {"url":url,"status":"skipped_too_large","bytes":len(data)}
            if not data: return {"url":url,"status":"empty"}
            suffix = safe_name(final_url, ctype)
            if suffix == ".pdf" and b"%PDF-" not in bytes(data[:1024]):
                return {
                    "url": url,
                    "final_url": final_url,
                    "title": candidate.get("title"),
                    "status": "download_error",
                    "content_type": ctype,
                    "bytes": len(data),
                    "error": "Phản hồi không có chữ ký PDF trong 1024 byte đầu; từ chối lưu nội dung HTML hoặc tệp không phải PDF.",
                }
            digest = hashlib.sha256(data).hexdigest()
            name = digest[:16] + suffix
            job_dir.mkdir(parents=True, exist_ok=True)
            path = job_dir / name
            path.write_bytes(data)
            extracted = extract_file(path, ctype)
            extracted_path = path.with_suffix(path.suffix + ".extracted.json")
            save(extracted_path, extracted)
            return {"url":url,"final_url":final_url,"title":candidate.get("title"),"host":urlparse(final_url).hostname,"status":"downloaded_needs_review","file":str(path.relative_to(ROOT)),"extracted_file":str(extracted_path.relative_to(ROOT)),"sha256":digest,"bytes":len(data),"content_type":ctype,"extraction":extracted.get("extraction"),"pages":extracted.get("page_count"),"downloaded_at":NOW,"source":"VietstockFinance document index","document_index":candidate.get("document_index"),"index_title":candidate.get("index_title") or candidate.get("title"),"period":candidate.get("period"),"report_year":candidate.get("report_year"),"published_at":candidate.get("published_at"),"last_update":candidate.get("last_update"),"file_ext":candidate.get("file_ext"),"document_id":candidate.get("document_id"),"report_type":candidate.get("report_type"),"official_filing_page":candidate.get("filing_page"),"filing_page_title":candidate.get("filing_page_title")}
    except Exception as exc:
        return {"url":url,"title":candidate.get("title"),"status":"download_error","error":str(exc)[:1000]}

DETAIL_SCHEMA_ACCEPTED = {
    "2026-10-22-bank-statement-qc-v3",
    "2026-10-22-bank-statement-qc-v4",
}
TERMINAL_REVIEW_STATUSES = {
    "manual_review_required",
    "manual_review_required_pending_extraction",
    "no_detail_fields_matched_pending_manual_review",
}


def ticker_has_source_page_gap(items: list[dict], max_pages: int = 14) -> bool:
    """Keep the current ticker active while any linked quarterly report lacks its first statement pages."""
    raw_root = (ROOT / "reports" / "raw").resolve()
    for item in items:
        if not isinstance(item, dict):
            continue
        source = str(item.get("source_archive_file") or "").strip()
        if not source:
            continue
        archive_path = (ROOT / source).resolve()
        try:
            archive_path.relative_to(raw_root)
        except ValueError:
            continue
        sidecar = Path(str(archive_path) + ".extracted.json")
        if not sidecar.is_file():
            return True
        try:
            payload = json.loads(sidecar.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return True
        if not isinstance(payload, dict):
            return True
        if payload.get("kind") == "pdf":
            documents = [payload]
        elif payload.get("kind") == "zip":
            documents = [
                member for member in payload.get("members", [])
                if isinstance(member, dict) and member.get("kind") == "pdf"
            ]
        else:
            documents = []
        for document in documents:
            try:
                page_count = int(document.get("page_count") or 0)
            except (TypeError, ValueError):
                page_count = 0
            if page_count <= 0:
                continue
            target_count = min(page_count, max_pages)
            pages = {
                int(page.get("page"))
                for page in document.get("pages", [])
                if isinstance(page, dict) and str(page.get("page", "")).isdigit()
            }
            if any(number not in pages for number in range(1, target_count + 1)):
                return True
    return False


def item_needs_vietstock_archive(symbol: str, item: object) -> bool:
    """True when a quarter has a direct consolidated Vietstock PDF URL but no stored archive."""
    if not isinstance(item, dict) or not re.fullmatch(r"20\d{2}-Q[1-4]", str(item.get("period") or "")):
        return False
    source_url = str(item.get("source_url") or "").strip()
    parsed = urlparse(source_url)
    host = (parsed.hostname or "").lower()
    path = unquote(parsed.path or "").lower()
    filename = Path(unquote(parsed.path or "")).name.lower()
    symbol = str(symbol or "").strip().upper()
    if (
        parsed.scheme != "https"
        or host not in {"static1.vietstock.vn", "static2.vietstock.vn"}
        or "/bctc/vn/" not in path
        or not path.endswith(".pdf")
        or not filename.startswith(symbol.lower() + "_")
    ):
        return False
    title = " ".join(str(item.get(key) or "") for key in ("report_type", "report_label", "source_archive_title", "title"))
    if re.search(r"công\s*ty\s*mẹ|cong\s*ty\s*me|riêng\s*lẻ|rieng\s*le|parent\s+company|separate", title + " " + filename, re.I):
        return False
    if not re.search(r"hợp\s*nhất|hop\s*nhat|consolidated", title + " " + filename, re.I):
        return False
    existing = str(item.get("source_archive_file") or "").strip()
    if not existing:
        return True
    archive = (ROOT / existing).resolve()
    try:
        archive.relative_to((ROOT / "reports" / "raw").resolve())
    except ValueError:
        return True
    return not archive.is_file()


def first_incomplete_symbol(securities: dict) -> str | None:
    for security in securities.get("securities", []):
        if not isinstance(security, dict) or security.get("enabled") is False:
            continue
        symbol = str(security.get("symbol") or "").strip().upper()
        if not symbol:
            continue
        path = ROOT / "data" / f"financials-{symbol}.json"
        if not path.exists():
            return symbol
        try:
            snapshot = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return symbol

        items = snapshot.get("items") or []
        # A report can be fully published on VietstockFinance after the last
        # catalog scan, even though this ticker is otherwise complete. Archive
        # that exact quarter PDF before skipping the completed ticker.
        if any(item_needs_vietstock_archive(symbol, item) for item in items):
            return symbol

        # The sync/OCR pipeline intentionally closes a ticker as
        # complete_needs_manual_review when automated extraction has reached a
        # terminal limit. Do not let the collector queue loop back to that
        # ticker forever just because its preserved source record still carries
        # manual_review_required_pending_extraction.
        if snapshot.get("bootstrap_status") in {"complete", "complete_needs_manual_review"}:
            continue
        if not items:
            try:
                attempts = int(snapshot.get("bootstrap_attempt_count") or 0)
            except (TypeError, ValueError):
                attempts = 0
            terminal = (
                attempts >= 3
                and snapshot.get("sync_status") in {
                    "fetch_failed", "no_quarterly_consolidated_records_returned",
                    "needs_manual_review_no_quarterly_data"
                }
            )
            if terminal:
                continue
            return symbol

        core_pending = any(
            isinstance(item, dict)
            and item.get("verification_status") in {
                "auto_pending_extraction",
                "manual_review_required_pending_extraction",
            }
            for item in items
        )
        details_pending = any(
            isinstance(item, dict)
            and re.fullmatch(r"20\d{2}-Q[1-4]", str(item.get("period") or ""))
            and item.get("total_assets") is not None
            and item.get("net_profit") is not None
            and not item.get("financial_line_items_status")
            for item in items
        )
        stale_details = any(
            isinstance(item, dict)
            and item.get("financial_line_items_status")
            and item.get("financial_line_items_status") not in TERMINAL_REVIEW_STATUSES
            and item.get("financial_line_items_schema_version") not in DETAIL_SCHEMA_ACCEPTED
            for item in items
        )
        statement_pages_pending = ticker_has_source_page_gap(items)
        if core_pending or details_pending or stale_details or statement_pages_pending:
            return symbol
    return None


def canonical_document_url(value: object) -> tuple[str, str] | None:
    try:
        parsed = urlparse(str(value or "").strip())
        if not parsed.scheme or not parsed.netloc:
            return None
        return (
            (parsed.hostname or "").casefold(),
            unquote(parsed.path).rstrip("/").casefold(),
        )
    except Exception:
        return None


FAILED_URL_RETRY_COOLDOWN_SECONDS = 30 * 60
FAILED_URL_RETRY_RESET_SECONDS = 24 * 60 * 60
FAILED_URL_MAX_ATTEMPTS = 3
CATALOG_REFRESH_INTERVAL_SECONDS = 6 * 60 * 60


def retry_url_key(value: object):
    """Normalize Vietstock URLs so a failed source can be retried safely."""
    return canonical_document_url(value) or str(value or "").strip()


def failed_url_retry_state(jobs: list[dict], symbol: str, now: datetime | None = None):
    """Return URL states with a rolling retry budget.

    A failed URL is retried after a cooldown, up to FAILED_URL_MAX_ATTEMPTS
    within a 24-hour window. Old failures expire instead of banning the source
    forever. Missing legacy timestamps are treated as expired metadata.
    """
    current = now or datetime.now(timezone.utc)
    now_epoch = current.timestamp()
    failures: dict[object, dict] = {}
    successful: set[object] = set()
    for prior_job in jobs:
        if str(prior_job.get("symbol") or "").strip().upper() != symbol:
            continue
        for record in prior_job.get("downloadedFiles") or []:
            if not isinstance(record, dict):
                continue
            status = str(record.get("status") or "")
            keys = {
                retry_url_key(record.get(field))
                for field in ("url", "final_url")
                if record.get(field)
            }
            for key in keys:
                if status == "downloaded_needs_review":
                    successful.add(key)
                elif status == "download_error":
                    stamp = record.get("attempted_at") or record.get("failed_at") or record.get("downloaded_at")
                    epoch = 0.0
                    if stamp:
                        try:
                            dt = datetime.fromisoformat(str(stamp).replace("Z", "+00:00"))
                            if dt.tzinfo is None:
                                dt = dt.replace(tzinfo=timezone.utc)
                            epoch = dt.timestamp()
                        except (TypeError, ValueError, OSError, OverflowError):
                            epoch = 0.0

                    state = failures.setdefault(key, {"attempts": 0, "last_attempt": 0.0})
                    # The retry budget is rolling: an old failed URL becomes
                    # eligible again after 24 hours, even if prior attempts hit
                    # the per-window maximum.
                    if epoch and now_epoch - epoch >= FAILED_URL_RETRY_RESET_SECONDS:
                        continue
                    if epoch:
                        state["attempts"] += 1
                        state["last_attempt"] = max(float(state.get("last_attempt") or 0), epoch)

    retryable = {
        key for key, state in failures.items()
        if key not in successful
        and int(state.get("attempts") or 0) < FAILED_URL_MAX_ATTEMPTS
        and (
            not float(state.get("last_attempt") or 0)
            or now_epoch - float(state["last_attempt"]) >= FAILED_URL_RETRY_COOLDOWN_SECONDS
        )
    }
    # A successful archive in any prior job always wins over a failure record.
    retryable.difference_update(successful)
    return successful, failures, retryable


def catalog_refresh_due(
    job: dict,
    has_pending_saved_candidates: bool = False,
    retryable_failed_urls: set | None = None,
    now: datetime | None = None,
) -> bool:
    """Throttle fresh catalog scans for jobs stuck after exhausting saved URLs."""
    if str(job.get("status") or "") not in {"needs_review", "failed"}:
        return False
    if has_pending_saved_candidates or retryable_failed_urls:
        return False
    stamp = job.get("lastCatalogRefreshAt")
    if not stamp:
        return True
    try:
        refreshed = datetime.fromisoformat(str(stamp).replace("Z", "+00:00"))
        if refreshed.tzinfo is None:
            refreshed = refreshed.replace(tzinfo=timezone.utc)
    except (TypeError, ValueError, OSError, OverflowError):
        return True
    current = now or datetime.now(timezone.utc)
    return (current - refreshed).total_seconds() >= CATALOG_REFRESH_INTERVAL_SECONDS


def link_quarterly_archive_sources(symbol: str, jobs: list[dict]) -> int:
    """Link only exact, consolidated quarterly Vietstock URLs to snapshot periods."""
    path = ROOT / "data" / f"financials-{symbol}.json"
    if not path.is_file():
        return 0
    try:
        snapshot = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return 0
    items = snapshot.get("items")
    if not isinstance(items, list):
        return 0
    by_period = {
        str(item.get("period")): item
        for item in items
        if isinstance(item, dict) and re.fullmatch(r"20\d{2}-Q[1-4]", str(item.get("period") or ""))
    }
    changed = 0
    for job in jobs:
        if str(job.get("symbol") or "").upper() != symbol:
            continue
        for record in job.get("downloadedFiles") or []:
            if not isinstance(record, dict) or record.get("status") != "downloaded_needs_review":
                continue
            source_period = str(record.get("period") or "")
            period = source_period
            half_year = re.fullmatch(r"(20\d{2})-H1", source_period)
            if half_year:
                # Reviewed half-year statements feed the Q2 snapshot only when
                # its exact source_url matches the archived document URL.
                period = half_year.group(1) + "-Q2"
            if not re.fullmatch(r"20\d{2}-Q[1-4]", period):
                continue
            title = " ".join(str(record.get(key) or "") for key in (
                "title", "index_title", "vietstock_title", "full_name"
            ))
            explicit_consolidated = re.search(r"hợp\s*nhất|hop\s*nhat|consolidated", title, re.I)
            explicit_separate = re.search(r"công\s*ty\s*mẹ|cong\s*ty\s*me|riêng\s*lẻ|rieng\s*le|parent\s+company|separate", title, re.I)
            official_financial_title = re.search(r"báo\s*cáo\s*tài\s*chính|bctc|financial\s+statements", title, re.I)
            # Neutral official VietstockFinance titles are permitted when the exact source URL
            # matches this saved quarter; explicit parent/separate reports are never linked.
            if explicit_separate or not (explicit_consolidated or official_financial_title):
                continue
            relative = str(record.get("file") or "").strip()
            if not relative:
                continue
            archive_path = (ROOT / relative).resolve()
            try:
                archive_path.relative_to((ROOT / "reports" / "raw").resolve())
            except ValueError:
                continue
            if not archive_path.is_file() or archive_path.suffix.lower() not in {".pdf", ".zip"}:
                continue
            item = by_period.get(period)
            if not item:
                continue
            report_url = item.get("source_url")
            download_url = record.get("url") or record.get("final_url")
            if not report_url or canonical_document_url(report_url) != canonical_document_url(download_url):
                continue
            existing = str(item.get("source_archive_file") or "").strip()
            existing_path = (ROOT / existing).resolve() if existing else None
            pending_full_report = (
                item.get("balance_sheet_extraction_status") == "pending_full_report_archive_and_extraction"
                and canonical_document_url(item.get("balance_sheet_source_url")) == canonical_document_url(download_url)
                and not (
                    isinstance(item.get("source_validation"), dict)
                    and item["source_validation"].get("full_report_present_in_archive") is True
                )
            )
            if existing_path and existing_path.is_file() and not pending_full_report:
                continue
            item["source_archive_file"] = relative
            item["source_archive_source"] = "VietstockFinance document index"
            item["source_archive_title"] = title
            item["source_archive_linked_at"] = NOW
            if pending_full_report:
                validation = item.get("source_validation")
                if not isinstance(validation, dict):
                    validation = {}
                validation["full_report_present_in_archive"] = True
                validation["full_report_extracted_via_direct_vietstock_pdf"] = False
                validation["full_report_archive_file"] = relative
                validation["note"] = "BCTC hợp nhất đầy đủ đã được lưu vào kho; chờ OCR/trích xuất bảng cân đối trước khi công bố số liệu."
                item["source_validation"] = validation
                item["core_financial_extraction_manual_terminal"] = False
                item["core_financial_extraction_evidence_attempt_count"] = 0
                item["core_financial_extraction_attempt_count"] = 0
                changed += 1
            changed += 1
    if changed:
        snapshot["items"] = items
        path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return changed


def main():
    global ALLOWED_DOMAINS
    sources = load(SOURCES_PATH, {"allowedDomains":[]})
    ALLOWED_DOMAINS = sorted(set(str(x).lower().strip(".") for x in sources.get("allowedDomains",[]) if x))
    if not ALLOWED_DOMAINS:
        print("No allowlisted domains configured. Add official domains to config/source-registry.json; worker will not download from unverified domains.", file=sys.stderr)
    queue = load(QUEUE_PATH, {"schemaVersion":1,"updatedAt":None,"jobs":[]})
    securities = load(SEC_PATH, {"securities":[]})
    jobs = queue.setdefault("jobs", [])
    initial_job_ids = {j.get("id") for j in jobs}
    known = {j.get("id") for j in jobs}
    # Refresh discovery weekly for every enabled security; avoid creating a new job every day.
    last_processed = {}
    for existing in jobs:
        symbol = existing.get("symbol")
        stamp = existing.get("processedAt") or existing.get("createdAt")
        if symbol and stamp and (symbol not in last_processed or stamp > last_processed[symbol]):
            last_processed[symbol] = stamp
    for sec in securities.get("securities",[]):
        symbol = str(sec.get("symbol","")).upper()
        last_stamp = last_processed.get(symbol)
        refresh_due = not last_stamp
        if last_stamp:
            try:
                age = datetime.now(timezone.utc) - datetime.fromisoformat(last_stamp.replace("Z","+00:00"))
                refresh_due = age.days >= 7
            except Exception:
                refresh_due = True
        if sec.get("enabled",True) and symbol and refresh_due:
            job_id = "auto-" + symbol.lower() + "-" + datetime.now(timezone.utc).strftime("%Y%m%d")
            if job_id not in known:
                jobs.append({"id":job_id,"symbol":symbol,"startYear":int(securities.get("defaultStartYear",2024)),"reportType":"both" if len(sec.get("reportTypes",[]))>1 else (sec.get("reportTypes") or ["consolidated"])[0],"status":"queued","createdAt":NOW,"attempts":0,"sourceCandidates":[],"downloadedFiles":[],"verifiedPeriods":[],"origin":"scheduled_registry"})
                known.add(job_id)
    summary = {"started_at":NOW,"jobs_processed":0,"downloaded":0,"needs_review":0,"errors":0,"allowlisted_domains":ALLOWED_DOMAINS}
    requested_symbol = os.getenv("REPORT_COLLECT_SYMBOL", "").strip().upper()
    enabled_symbols = {
        str(sec.get("symbol") or "").strip().upper()
        for sec in securities.get("securities", [])
        if sec.get("enabled", True) and sec.get("symbol")
    }
    if requested_symbol and requested_symbol not in enabled_symbols:
        print(f"Requested collector ticker is not enabled: {requested_symbol}", file=sys.stderr)
        return 2
    active_symbol = requested_symbol or first_incomplete_symbol(securities)

    # If a quarterly snapshot explicitly points to a full consolidated report
    # that is still missing from the archive, create a targeted one-document
    # repair job. This prevents a previously exhausted general catalog scan from
    # silently holding the ticker forever without retrying the exact source URL.
    if active_symbol:
        snapshot_path = ROOT / "data" / f"financials-{active_symbol}.json"
        try:
            snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            snapshot = {}
        for item in snapshot.get("items", []) if isinstance(snapshot, dict) else []:
            if not isinstance(item, dict):
                continue
            balance_sheet_pending = item.get("balance_sheet_extraction_status") == "pending_full_report_archive_and_extraction"
            source_archive_pending = item_needs_vietstock_archive(active_symbol, item)
            if not balance_sheet_pending and not source_archive_pending:
                continue
            source_url = str(
                (item.get("balance_sheet_source_url") if balance_sheet_pending else item.get("source_url")) or ""
            ).strip()
            period = str(item.get("period") or "").strip().upper()
            if not source_url or not re.fullmatch(r"20\d{2}-Q[1-4]", period):
                continue
            validation = item.get("source_validation")
            if (
                isinstance(validation, dict)
                and validation.get("full_report_present_in_archive") is True
                and not source_archive_pending
            ):
                continue
            repair_id = f"repair-{active_symbol.lower()}-{period.lower().replace('-', '')}-full-report"
            existing = next((job for job in jobs if job.get("id") == repair_id), None)
            if existing is None:
                year = int(period[:4])
                title = f"Báo cáo tài chính hợp nhất {active_symbol} {period}"
                existing = {
                    "id": repair_id,
                    "symbol": active_symbol,
                    "startYear": year,
                    "reportType": "consolidated",
                    "status": "queued",
                    "createdAt": NOW,
                    "attempts": 0,
                    "sourceCandidates": [],
                    "requestedUrls": [{
                        "url": source_url,
                        "title": title,
                        "period": period,
                        "report_year": year,
                    }],
                    "downloadedFiles": [],
                    "verifiedPeriods": [],
                    "origin": "missing_full_report_dependency",
                    "priority": 1,
                }
                jobs.append(existing)
                known.add(repair_id)
                print(json.dumps({
                    "event": "targeted_full_report_repair_queued",
                    "symbol": active_symbol,
                    "period": period,
                    "url": source_url,
                    "job_id": repair_id,
                }, ensure_ascii=False))
            elif (
                existing.get("origin") == "missing_full_report_dependency"
                and existing.get("status") == "failed"
                and int(existing.get("attempts", 0)) < 5
            ):
                existing["status"] = "queued"
                existing["updatedAt"] = NOW
                print(json.dumps({
                    "event": "targeted_full_report_repair_requeued",
                    "symbol": active_symbol,
                    "period": period,
                    "url": source_url,
                    "job_id": repair_id,
                    "attempts": int(existing.get("attempts", 0)),
                }, ensure_ascii=False))
            break

    print(json.dumps({
        "event": "single_ticker_queue_head",
        "symbol": active_symbol,
        "selection": "explicit_handoff" if requested_symbol else "first_incomplete_ticker",
    }, ensure_ascii=False))
    # A ticker remains at the head of the queue while its discovered source URLs
    # are being archived. Each run downloads at most MAX_PER_JOB documents; only
    # after the remaining candidate list is exhausted does the worker move on.
    jobs.sort(key=lambda job: (
        0 if job.get("resumeDownload") else (1 if job.get("status") == "queued" else 2),
        int(job.get("priority", 50)),
        0 if str(job.get("symbol", "")).upper() == "MBB" else 1,
        job.get("createdAt", ""),
    ))
    processed_symbols = set()
    for job in jobs:
        if active_symbol and str(job.get("symbol") or "").upper() != active_symbol:
            continue
        if summary["jobs_processed"] >= MAX_JOBS_PER_RUN:
            break
        symbol = str(job.get("symbol") or "").strip().upper()
        successful_url_keys, failed_url_states, retryable_failed_urls = failed_url_retry_state(jobs, symbol)
        stored_candidates = [
            candidate for candidate in (job.get("sourceCandidates") or [])
            if isinstance(candidate, dict) and candidate.get("url")
        ]
        attempted_candidate_keys = {
            retry_url_key(url)
            for url in (job.get("attemptedUrls") or [])
            if url
        }
        has_pending_saved_candidates = any(
            retry_url_key(candidate.get("url"))
            and retry_url_key(candidate.get("url")) not in successful_url_keys
            and (
                retry_url_key(candidate.get("url")) not in attempted_candidate_keys
                or retry_url_key(candidate.get("url")) in retryable_failed_urls
            )
            for candidate in stored_candidates
        )
        # Resume unattempted or retryable checkpointed URLs first. If every saved
        # URL is exhausted, periodically rescan VietstockFinance for new/updated
        # document links rather than leaving the ticker permanently stranded.
        resume_pending_candidates = (
            job.get("status") == "needs_review" and has_pending_saved_candidates
        )
        retry_failed_only = job.get("status") == "needs_review" and bool(retryable_failed_urls)
        refresh_catalog = catalog_refresh_due(
            job,
            has_pending_saved_candidates=has_pending_saved_candidates,
            retryable_failed_urls=retryable_failed_urls,
        )
        if (
            job.get("status") not in {"queued", "failed"}
            and not retry_failed_only
            and not resume_pending_candidates
            and not refresh_catalog
        ):
            continue
        if (
            job.get("status") == "failed"
            and int(job.get("attempts", 0)) >= 5
            and job.get("origin") != "missing_full_report_dependency"
            and not retryable_failed_urls
            and not refresh_catalog
        ):
            continue

        previous_files = list(job.get("downloadedFiles") or [])
        attempted_urls = set(str(url) for url in (job.get("attemptedUrls") or []) if url)
        # A resume job already carries the remaining catalog URLs from the previous
        # successful/partial batch. Reusing these URLs avoids re-querying every
        # VietstockFinance year/page before each two-file download batch.
        resume_saved_candidates = (
            (
                job.get("status") == "queued" and job.get("resumeDownload") is True
            )
            or resume_pending_candidates
        ) and isinstance(job.get("sourceCandidates"), list) and bool(job.get("sourceCandidates"))
        saved_candidates = [
            candidate for candidate in (job.get("sourceCandidates") or [])
            if isinstance(candidate, dict) and candidate.get("url")
        ] if resume_saved_candidates else []
        processed_symbols.add(symbol)
        job["status"] = "searching"
        job["updatedAt"] = NOW
        job["attempts"] = int(job.get("attempts", 0)) + 1
        if saved_candidates:
            candidates = saved_candidates
            print(json.dumps({
                "event": "resume_saved_source_candidates",
                "symbol": symbol,
                "candidate_count": len(candidates),
                "note": "Reusing checkpointed VietstockFinance URLs for this ticker; no full catalog rescan for the continuation batch."
            }, ensure_ascii=False))
        else:
            candidates = discover_for_job(job, ALLOWED_DOMAINS)
            if not (isinstance(job.get("requestedUrls"), list) and job.get("requestedUrls")):
                job["lastCatalogRefreshAt"] = NOW
                if refresh_catalog:
                    print(json.dumps({
                        "event": "stuck_job_catalog_refresh",
                        "symbol": symbol,
                        "job_id": job.get("id"),
                        "candidate_count": len(candidates),
                        "refresh_interval_hours": CATALOG_REFRESH_INTERVAL_SECONDS // 3600,
                    }, ensure_ascii=False))

        # Avoid re-downloading successful archives, but allow bounded retries for
        # failed URLs after the cooldown. attemptedUrls alone must not permanently
        # hide a transient Vietstock 404/5xx or rate-limit response.
        already_downloaded_urls = set(successful_url_keys)
        for attempted_url in attempted_urls:
            key = retry_url_key(attempted_url)
            if key not in retryable_failed_urls:
                already_downloaded_urls.add(key)
        for failed_url in failed_url_states:
            if failed_url not in retryable_failed_urls:
                already_downloaded_urls.add(failed_url)
        # A targeted repair exists specifically because the linked full report is
        # still missing. Permit a bounded retry of that exact URL even if older
        # general catalog jobs exhausted their retries; successful archives still win.
        if job.get("origin") == "missing_full_report_dependency":
            for request in job.get("requestedUrls") or []:
                if not isinstance(request, dict):
                    continue
                requested_url = str(request.get("url") or "")
                requested_key = retry_url_key(requested_url)
                requested_period = str(request.get("period") or "").upper()
                if not requested_key:
                    continue
                dependency_still_pending = any(
                    isinstance(item, dict)
                    and str(item.get("period") or "").upper() == requested_period
                    and (
                        (
                            item.get("balance_sheet_extraction_status") == "pending_full_report_archive_and_extraction"
                            and str(item.get("balance_sheet_source_url") or "").strip() == requested_url
                        )
                        or (
                            str(item.get("source_url") or "").strip() == requested_url
                            and item_needs_vietstock_archive(active_symbol, item)
                        )
                    )
                    and not (
                        isinstance(item.get("source_validation"), dict)
                        and item["source_validation"].get("full_report_present_in_archive") is True
                        and not item_needs_vietstock_archive(active_symbol, item)
                    )
                    for item in (snapshot.get("items") or [])
                ) if isinstance(snapshot, dict) else False
                if dependency_still_pending:
                    # The snapshot itself says this exact report is still missing.
                    # Do not let stale queue success metadata suppress the repair.
                    successful_url_keys.discard(requested_key)
                    already_downloaded_urls.discard(requested_key)

        before_dedup = len(candidates)
        filtered_candidates = []
        seen_candidate_urls = set()
        for candidate in candidates:
            key = retry_url_key(candidate.get("url"))
            if not key or key in already_downloaded_urls or key in seen_candidate_urls:
                continue
            filtered_candidates.append(candidate)
            seen_candidate_urls.add(key)
        candidates = filtered_candidates
        job["sourceCandidates"] = candidates[:80]
        print(json.dumps({
            "event": "deduplicate_previous_downloads",
            "symbol": job.get("symbol"),
            "candidate_count_before": before_dedup,
            "already_downloaded_urls": len(already_downloaded_urls),
            "candidate_count_after": len(candidates),
        }, ensure_ascii=False))

        if not candidates:
            if any(item.get("status") == "downloaded_needs_review" for item in previous_files):
                job["downloadedFiles"] = previous_files
                job["status"] = "needs_review"
                job["resumeDownload"] = False
                job["lastError"] = None
                job["processedAt"] = NOW
                summary["jobs_processed"] += 1
                summary["needs_review"] += 1
                print(json.dumps({
                    "job": job["id"], "symbol": job["symbol"],
                    "status": job["status"], "event": "ticker_archive_batch_complete",
                    "downloaded_total": len(previous_files),
                }, ensure_ascii=False))
                continue
            job["downloadedFiles"] = previous_files
            job["status"] = "failed"
            job["resumeDownload"] = False
            job["lastError"] = "Không tìm được thêm URL báo cáo có thể tải trong miền cho phép."
            job["processedAt"] = NOW
            summary["errors"] += 1
            summary["jobs_processed"] += 1
            continue

        job_dir = OUT / job["symbol"] / job["id"]
        results = []
        attempted_this_run = []
        successes_this_run = 0
        for candidate in candidates:
            # Bound attempted downloads, not just successful downloads. Broken
            # source links must not cause the worker to sweep the whole catalog.
            if len(attempted_this_run) >= MAX_PER_JOB:
                break
            result = download_candidate(candidate, job_dir)
            candidate_url = str(candidate.get("url", "") or "")
            candidate_key = retry_url_key(candidate_url)
            if result.get("status") == "download_error":
                previous_failure = failed_url_states.get(candidate_key, {})
                result["attempted_at"] = NOW
                result["attempt_number"] = int(previous_failure.get("attempts") or 0) + 1
                failed_url_states[candidate_key] = {
                    "attempts": result["attempt_number"],
                    "last_attempt": datetime.now(timezone.utc).timestamp(),
                }
                retryable_failed_urls.discard(candidate_key)
            results.append(result)
            attempted_this_run.append(candidate_url)
            if result.get("status") == "downloaded_needs_review":
                summary["downloaded"] += 1
                successes_this_run += 1
            if result.get("status") == "download_error":
                summary["errors"] += 1

        attempted_urls.update(url for url in attempted_this_run if url)
        job["attemptedUrls"] = sorted(attempted_urls)
        combined_files = previous_files + results
        job["downloadedFiles"] = combined_files
        more_candidates = len(attempted_this_run) < len(candidates)
        downloaded_any = any(item.get("status") == "downloaded_needs_review" for item in combined_files)

        if more_candidates:
            # Keep the same ticker at the front of the queue, even if this
            # batch contained only failed URLs; the next run will continue on
            # the next unattempted candidates without sweeping the full catalog.
            job["status"] = "queued"
            job["resumeDownload"] = True
            job["lastError"] = None if downloaded_any else (
                "Một lượt tải chỉ thử tối đa hai URL; sẽ tiếp tục các URL còn lại ở lượt sau."
            )
        elif downloaded_any:
            job["status"] = "needs_review"
            job["resumeDownload"] = False
            job["lastError"] = None
        else:
            job["status"] = "failed"
            job["resumeDownload"] = False
            job["lastError"] = "Không tải được tài liệu từ miền đã cho phép; kiểm tra sourceCandidates và cấu hình allowlist."

        job["processedAt"] = NOW
        summary["jobs_processed"] += 1
        if job["status"] == "needs_review":
            summary["needs_review"] += 1
        print(json.dumps({
            "job": job["id"], "symbol": job["symbol"], "status": job["status"],
            "candidates": before_dedup,
            "downloaded_this_run": successes_this_run,
            "downloaded_total": sum(1 for item in combined_files if item.get("status") == "downloaded_needs_review"),
            "more_candidates": more_candidates,
        }, ensure_ascii=False))
    active_job = next(
        (
            job for job in jobs
            if active_symbol and str(job.get("symbol") or "").strip().upper() == active_symbol
            and (job.get("status") in {"queued", "failed", "needs_review"} or job.get("resumeDownload"))
        ),
        None,
    )
    summary["active_symbol"] = active_symbol
    summary["active_job_id"] = active_job.get("id") if active_job else None
    summary["active_job_status"] = active_job.get("status") if active_job else None
    summary["remaining_source_candidates"] = len(active_job.get("sourceCandidates") or []) if active_job else 0
    # Chain only when the previous batch actually ran and explicitly left more
    # candidate URLs for the same ticker. A cooldown/no-progress run must stop.
    summary["continuation_needed"] = bool(
        summary["jobs_processed"] > 0
        and active_job
        and active_job.get("status") == "queued"
        and active_job.get("resumeDownload") is True
    )
    # Also backfill exact source-to-archive links for the current queue-head
    # ticker even when its prior queue jobs are already marked needs_review.
    # This lets the OCR/enrichment workflow finish this ticker before moving on.
    linked_count = link_quarterly_archive_sources(active_symbol, jobs) if active_symbol else 0
    if linked_count:
        print(json.dumps({"event":"quarterly_archive_sources_linked","symbol":active_symbol,"count":linked_count},ensure_ascii=False))
    queue_changed = summary["jobs_processed"] > 0 or linked_count > 0 or any(job.get("id") not in initial_job_ids for job in jobs)
    if queue_changed:
        queue["updatedAt"]=NOW
        save(QUEUE_PATH, queue)
    else:
        print(json.dumps({"event":"no_queue_change","active_symbol":active_symbol,"note":"Keep the ticker on hold rather than moving to another symbol."},ensure_ascii=False))
    save(ROOT / "reports/worker-last-run.json", summary)
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__ == "__main__":
    main()
